from openpyxl import load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter
from pathlib import Path

file_path = Path(r"C:\Users\Administrator\.openclaw\workspace\Stonewater_Karjat_Quote.xlsx")
wb = load_workbook(file_path)
ws = wb["All_In_One"]

ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
name = "Table_All_In_One"
if name in ws.tables:
    del ws.tables[name]

table = Table(displayName=name, ref=ref)
style = TableStyleInfo(
    name="TableStyleMedium9",
    showFirstColumn=False,
    showLastColumn=False,
    showRowStripes=True,
    showColumnStripes=False,
)
table.tableStyleInfo = style
ws.add_table(table)

for col in ws.columns:
    max_len = 0
    col_letter = get_column_letter(col[0].column)
    for cell in col:
        val = "" if cell.value is None else str(cell.value)
        max_len = max(max_len, len(val))
    ws.column_dimensions[col_letter].width = min(max_len + 2, 80)

wb.save(file_path)
print(str(file_path))
