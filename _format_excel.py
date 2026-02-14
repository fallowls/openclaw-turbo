from openpyxl import load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter
from pathlib import Path

file_path = Path(r"C:\Users\Administrator\.openclaw\workspace\Stonewater_Karjat_Quote.xlsx")
wb = load_workbook(file_path)

for ws in wb.worksheets:
    # Determine data range
    if ws.max_row < 2 or ws.max_column < 1:
        continue

    # Add table over used range
    ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    table_name = f"Table_{ws.title.replace(' ', '_')}"
    # Remove existing table with same name if any
    if table_name in ws.tables:
        del ws.tables[table_name]

    tab = Table(displayName=table_name, ref=ref)
    style = TableStyleInfo(
        name="TableStyleMedium9",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    tab.tableStyleInfo = style
    ws.add_table(tab)

    # Auto-fit column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                max_len = max(max_len, len(str(cell.value)) if cell.value is not None else 0)
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max_len + 2, 60)

wb.save(file_path)
print(str(file_path))
