from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font
from pathlib import Path

file_path = Path(r"C:\Users\Administrator\.openclaw\workspace\Stonewater_Karjat_Quote.xlsx")
wb = load_workbook(file_path)
ws = wb["All_In_One"]

header_fill = PatternFill("solid", fgColor="1F2937")
header_font = Font(color="FFFFFF", bold=True)
alt_fill = PatternFill("solid", fgColor="F3F4F6")

# Apply header style
for cell in ws[1]:
    cell.fill = header_fill
    cell.font = header_font

# Apply alternating row fill (starting from row 2)
for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
    if row[0].row % 2 == 0:
        for cell in row:
            cell.fill = alt_fill

wb.save(file_path)
print(str(file_path))
