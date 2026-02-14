from openpyxl import load_workbook, Workbook
from pathlib import Path

src = Path(r"C:\Users\Administrator\.openclaw\workspace\Stonewater_Karjat_Quote.xlsx")
dst = Path(r"C:\Users\Administrator\.openclaw\workspace\Stonewater_Karjat_Quote_AllInOne.xlsx")

wb = load_workbook(src)
ws = wb["All_In_One"]

new_wb = Workbook()
new_ws = new_wb.active
new_ws.title = "All_In_One"

for row in ws.iter_rows():
    new_ws.append([cell.value for cell in row])

# Copy column widths
for col, dim in ws.column_dimensions.items():
    new_ws.column_dimensions[col].width = dim.width

# Copy styles (header + alternating fills)
from copy import copy
for row in ws.iter_rows():
    for cell in row:
        new_cell = new_ws.cell(row=cell.row, column=cell.column)
        if cell.has_style:
            new_cell.font = copy(cell.font)
            new_cell.fill = copy(cell.fill)
            new_cell.border = copy(cell.border)
            new_cell.alignment = copy(cell.alignment)
            new_cell.number_format = copy(cell.number_format)

new_wb.save(dst)
print(str(dst))
